from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "reports" / "lumniverse_agent_efficiency_benchmark_data.json"
OUT_DOCX = ROOT / "reports" / "lumniverse_agent_efficiency_benchmark.docx"

NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x68, 0x73, 0x81)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
TEAL_FILL = "E8F4F2"
GOLD_FILL = "FFF6E3"


def money(value: float | None, decimals: int = 0) -> str:
    if value is None:
        return "Unavailable"
    return f"${value:,.{decimals}f}"


def num(value: float | int | None, decimals: int = 0) -> str:
    if value is None:
        return "Unavailable"
    return f"{value:,.{decimals}f}"


def pct(value: float | None, decimals: int = 1) -> str:
    if value is None:
        return "Unavailable"
    return f"{100.0 * value:.{decimals}f}%"


def pct_raw(value: float | None, decimals: int = 1) -> str:
    if value is None:
        return "Unavailable"
    return f"{value:.{decimals}f}%"


def hours(seconds: float | None) -> str:
    if seconds is None:
        return "Unavailable"
    return f"{seconds / 3600:,.1f} hours"


def set_run_font(run, size: float | None = None, color: RGBColor | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def remove_children(parent, tag: str) -> None:
    qtag = qn(tag)
    for child in list(parent):
        if child.tag == qtag:
            parent.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    remove_children(tc_pr, "w:shd")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    remove_children(tc_pr, "w:tcMar")
    tc_mar = OxmlElement("w:tcMar")
    for key, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    remove_children(tc_pr, "w:tcW")
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    remove_children(tbl_pr, "w:tblW")
    remove_children(tbl_pr, "w:tblInd")
    remove_children(tbl_pr, "w:tblLayout")

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)

    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    remove_children(tbl, "w:tblGrid")
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def paragraph_border_bottom(paragraph, color: str = "D9E2EC", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    remove_children(p_pr, "w:pBdr")
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_field_run(paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.color.rgb = GRAY
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(2)
    run = header.add_run("Agent Efficiency Benchmark Report")
    set_run_font(run, size=9, color=GRAY)
    paragraph_border_bottom(header, color="D9E2EC", size="4")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    set_run_font(run, size=9, color=GRAY)
    add_field_run(footer, "PAGE", "1")


def add_para(doc: Document, text: str = "", size: float | None = None, color: RGBColor | None = None, bold: bool | None = None, italic: bool | None = None, align: int | None = None, before: float | None = None, after: float | None = None, keep: bool = False):
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p


def add_rich_para(doc: Document, parts: Iterable[tuple[str, bool]]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, bold=bold)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_callout(doc: Document, label: str, body: str, fill: str = CALLOUT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9120])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label)
    set_run_font(run, size=10.5, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run = p2.add_run(body)
    set_run_font(run, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    run = p.add_run(text)
    set_run_font(run, size=9, color=GRAY, italic=True)


def add_figure(doc: Document, rel_path: str, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(ROOT / rel_path), width=Inches(6.2))
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], header_fill: str = LIGHT_FILL) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_header_row(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=9.5, color=NAVY, bold=True)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_metric_table(doc: Document, rows: list[tuple[str, str, str]], fill: str = TEAL_FILL) -> None:
    add_table(
        doc,
        ["Metric", "Value", "Why it matters"],
        [[a, b, c] for a, b, c in rows],
        [2300, 1900, 5000],
        header_fill=fill,
    )


def title_page(doc: Document, data: dict) -> None:
    add_para(doc, "BENCHMARK INSIGHTS REPORT", size=10, color=GRAY, bold=True, after=18, keep=True)
    p = add_para(doc, "Why Agent Benchmarks Need Trajectory Intelligence", size=28, color=NAVY, bold=True, after=8, keep=True)
    p.paragraph_format.line_spacing = 1.0
    add_para(
        doc,
        "Evidence from Terminal-Bench, SWE-agent trajectories, and AFWB Foundation",
        size=14,
        color=DARK_BLUE,
        after=18,
    )
    add_para(
        doc,
        "Pass rate is useful, but incomplete. The benchmark evidence shows that agent teams also need trajectory, cost, waste, root-cause, and scaffold-level analysis.",
        size=12,
        color=RGBColor(0x26, 0x32, 0x38),
        after=16,
    )
    add_table(
        doc,
        ["Dataset family", "Coverage used", "Primary use in this report"],
        [
            ["Terminal-Bench", "52,104 runs; cost and duration available", "FinOps, failure cost, pass/fail, trajectory waste where available"],
            ["SWE-agent", "80,036 runs; full trajectories; no cost/duration", "Edit/test loop behavior, repeat behavior, scaffold and trajectory divergence"],
            ["AFWB Foundation", "4 deterministic RCA scenarios", "Controlled root-cause and remediation examples"],
        ],
        [2100, 3100, 4000],
        header_fill=LIGHT_FILL,
    )
    add_callout(
        doc,
        "Reading note",
        "Dollar figures in this report come only from Terminal-Bench because that dataset has cost fields. SWE-agent is used for trajectory behavior only, and AFWB is used as deterministic RCA proof.",
        fill=GOLD_FILL,
    )
    doc.add_page_break()


def executive_summary(doc: Document, data: dict) -> None:
    t = data["terminal_overview"]
    s = data["swe_overview_direct"]
    e = data["terminal_early_stop_max"]
    add_heading(doc, "Executive Summary", 1)
    add_callout(
        doc,
        "Core finding",
        "Agent quality cannot be evaluated by final pass/fail alone. Failed runs are expensive, successful runs can still be wasteful, and the same model can behave very differently under different agent scaffolds.",
        fill=TEAL_FILL,
    )
    add_metric_table(
        doc,
        [
            ["Terminal-Bench pass rate", pct(t["pass_rate"]), "A useful outcome metric, but not enough to explain cost or failure mode."],
            ["Terminal-Bench trajectory coverage", pct(t["trajectory_coverage"]), "About one-third of runs lack usable trajectories, which limits RCA."],
            ["SWE-agent pass rate", pct(s["pass_rate"]), "The full processed SWE-agent data shows harder software-repair trajectories."],
            ["SWE-agent cost and duration", "Unavailable", "This prevents honest FinOps claims on SWE-agent alone."],
            ["Conservative early-stop savings", f"{money(e['estimated_cost_saved_usd'], 2)} and {hours(e['estimated_duration_saved_seconds'])}", "Counterfactual historical estimate from Terminal-Bench."],
        ],
    )
    add_figure(
        doc,
        data["figures"]["benchmark_coverage"],
        "Figure 1. Pass rate and trajectory coverage answer different benchmark questions.",
    )
    add_para(
        doc,
        "The report keeps dataset families separate. Terminal-Bench supports cost and duration analysis; SWE-agent supports trajectory-loop analysis; AFWB supports controlled RCA and remediation examples.",
    )
    doc.add_page_break()


def methodology(doc: Document, data: dict) -> None:
    tq = data["terminal_quality"]
    sq = data["swe_quality"]
    add_heading(doc, "Methodology And Caveats", 1)
    add_para(
        doc,
        "The analysis uses local processed artifacts from the benchmark repository. It does not merge dataset families into a universal leaderboard because each dataset has different coverage, metadata, and measurement semantics.",
    )
    add_table(
        doc,
        ["Dataset", "Source revision / cutoff", "What is measured", "Important caveat"],
        [
            [
                "Terminal-Bench",
                f"{tq['source_revision'][:12]}...; cutoff {tq['snapshot_cutoff']}",
                "Outcome, trajectories when available, cost, duration, tokens",
                "Only 66.1% of runs have usable trajectory steps.",
            ],
            [
                "SWE-agent",
                f"{sq['source_revision'][:12]}...; source last modified {sq['source_last_modified']}",
                "Outcome and full normalized trajectories",
                "Cost, duration, and token fields are unavailable in the processed source.",
            ],
            [
                "AFWB Foundation",
                "Local deterministic foundation suite",
                "Known root cause, first observable failure, remediation, avoidable usage",
                "Controlled RCA proof layer, not public-scale benchmark evidence.",
            ],
        ],
        [1900, 2500, 2900, 2100],
    )
    add_callout(
        doc,
        "Counterfactual note",
        "Early-stopping savings are historical counterfactual estimates: they estimate what could have been saved on these recorded trajectories if a stop rule had been applied.",
        fill=GOLD_FILL,
    )


def q1_cost_failed_runs(doc: Document, data: dict) -> None:
    rows = {row["success"]: row for row in data["terminal_outcome_cost"]}
    failed = rows[False]
    passed = rows[True]
    add_heading(doc, "1. What Is The Cost Of Failed AI Agent Runs?", 1)
    add_callout(
        doc,
        "Direct answer",
        f"Terminal-Bench failed runs consumed {money(failed['total_cost_usd'])} in observed cost, about {pct(data['derived']['terminal_failed_cost_share'])} of all observed run cost. Failed runs were {data['derived']['terminal_avg_failed_to_success_cost_ratio']:.1f}x more expensive on average than successful runs: {money(failed['avg_cost_usd'], 3)} vs. {money(passed['avg_cost_usd'], 3)}.",
    )
    add_figure(doc, data["figures"]["cost_outcome"], "Figure 2. Failed runs dominate observed Terminal-Bench spend.")
    add_metric_table(
        doc,
        [
            ["Failed runs", f"{num(failed['runs'])} runs; {money(failed['total_cost_usd'])} observed cost", "Failure burns budget before the final negative outcome."],
            ["Successful runs", f"{num(passed['runs'])} runs; {money(passed['total_cost_usd'])} observed cost", "Successful outcomes still need efficiency analysis."],
            ["Median duration", f"{num(failed['median_duration_seconds'])} sec failed vs. {num(passed['median_duration_seconds'])} sec successful", "Failed runs tend to last longer."],
        ],
    )
    add_para(doc, "Interpretation: cost-of-failure analysis should be a first-class benchmark dimension. A pass-rate-only table hides where the budget actually goes.")


def q2_successful_waste(doc: Document, data: dict) -> None:
    row = data["terminal_success_waste"]
    add_heading(doc, "2. Why Can Successful Agent Runs Still Be Wasteful?", 1)
    add_callout(
        doc,
        "Direct answer",
        f"Successful Terminal-Bench trajectory runs still contained {num(row['suspected_wasted_actions'])} suspected wasted actions, {num(row['error_events'])} error events, and {num(row['adjacent_repeats'])} adjacent repeats. Estimated waste inside successful runs was {money(row['estimated_wasted_cost_usd'], 2)} and {hours(row['estimated_wasted_duration_seconds'])}.",
    )
    add_figure(doc, data["figures"]["success_waste"], "Figure 3. Passing runs can still contain repeated actions and error recovery overhead.")
    add_para(
        doc,
        "Interpretation: a passing run can still be brittle, slow, or expensive. The useful question is not only whether the agent eventually succeeded, but how much avoidable work it did on the way.",
    )


def q3_repeated_actions(doc: Document, data: dict) -> None:
    strict = data["terminal_repeated_failed_agents"][:6]
    repeat = data["terminal_repeat_rate_agents"][:5]
    add_heading(doc, "3. Which Coding Agents Repeat Failed Actions Most Often?", 1)
    add_callout(
        doc,
        "Direct answer",
        "Strict repeated-failed-approach RCA is concentrated in a small number of agents. In rate terms, judy is highest at 9.55% of trajectory runs, but with a small sample. In count terms, mini-swe-agent has 55 runs and terminus-2 has 42 runs.",
    )
    add_figure(doc, data["figures"]["repeated_actions"], "Figure 4. Strict repeated-failed-approach rate among agents with at least 100 trajectory runs.")
    add_table(
        doc,
        ["Agent", "Trajectory runs", "Strict repeated-failed runs", "Rate", "Repeat behavior note"],
        [
            [
                row["agent"],
                num(row["runs"]),
                num(row["repeated_failed_rca_runs"]),
                pct_raw(row["repeated_failed_rca_rate_pct"]),
                "High strict RCA rate" if row["agent"] == "judy" else "High volume" if row["agent"] in {"mini-swe-agent", "terminus-2"} else "Observed pattern",
            ]
            for row in strict
        ],
        [1500, 1700, 2100, 1300, 2700],
    )
    add_para(
        doc,
        f"For general repeat behavior, not only strict failed repeats, the highest mean repeat rates were {repeat[0]['agent']} at {pct(repeat[0]['mean_repeat_rate'])} and {repeat[1]['agent']} at {pct(repeat[1]['mean_repeat_rate'])}.",
    )
    add_para(doc, "Interpretation: repeated actions are invisible from final outcome alone. They require step-level trajectory analysis.")


def q4_divergence(doc: Document, data: dict) -> None:
    t = {row["success"]: row for row in data["terminal_trajectory_outcome"]}
    s = {row["success"]: row for row in data["swe_outcome_direct"]}
    add_heading(doc, "4. Where Do Successful And Failed Trajectories Diverge?", 1)
    add_callout(
        doc,
        "Direct answer",
        f"Terminal-Bench failed trajectory runs had higher medians than successful runs: {num(t[False]['median_duration_seconds'])} sec vs. {num(t[True]['median_duration_seconds'])} sec, {num(t[False]['median_steps'])} vs. {num(t[True]['median_steps'])} steps, and {num(t[False]['median_tool_calls'])} vs. {num(t[True]['median_tool_calls'])} tool calls. SWE-agent failed runs also used more steps and tool calls, with repeat rate {pct(s[False]['mean_repeat_rate'])} vs. {pct(s[True]['mean_repeat_rate'])}.",
    )
    add_figure(doc, data["figures"]["divergence"], "Figure 5. Failed trajectories show heavier execution patterns before final outcome.")
    add_table(
        doc,
        ["Dataset", "Failed trajectory profile", "Successful trajectory profile", "Most useful signal"],
        [
            [
                "Terminal-Bench",
                f"{num(t[False]['median_duration_seconds'])} sec; {num(t[False]['median_steps'])} steps; {num(t[False]['median_tool_calls'])} tool calls",
                f"{num(t[True]['median_duration_seconds'])} sec; {num(t[True]['median_steps'])} steps; {num(t[True]['median_tool_calls'])} tool calls",
                f"Error rate {pct(t[False]['mean_error_event_rate'])} failed vs. {pct(t[True]['mean_error_event_rate'])} successful",
            ],
            [
                "SWE-agent",
                f"{num(s[False]['median_steps'])} steps; {num(s[False]['median_tool_calls'])} tool calls",
                f"{num(s[True]['median_steps'])} steps; {num(s[True]['median_tool_calls'])} tool calls",
                f"Repeat rate {pct(s[False]['mean_repeat_rate'])} failed vs. {pct(s[True]['mean_repeat_rate'])} successful",
            ],
        ],
        [1700, 2700, 2700, 2200],
    )
    add_para(doc, "Interpretation: trajectory signals can reveal a run going bad before the final answer is submitted.")


def q5_early_stop(doc: Document, data: dict) -> None:
    e = data["terminal_early_stop_max"]
    policies = data["terminal_early_stop_policy"]
    add_heading(doc, "5. How Much Could Early Stopping Save?", 1)
    add_callout(
        doc,
        "Direct answer",
        f"On Terminal-Bench, conservative early stopping could have saved about {money(e['estimated_cost_saved_usd'], 2)}, {num(e['events_saved'])} trajectory events, and {hours(e['estimated_duration_saved_seconds'])}. It would trigger on {num(e['conservative_triggered_runs'])} runs, or {pct_raw(e['conservative_trigger_rate_pct'])} of all runs.",
    )
    add_figure(doc, data["figures"]["early_stop"], "Figure 6. Conservative early-stop savings by individual policy.")
    add_table(
        doc,
        ["Policy", "Conservative triggers", "Saved events", "Estimated saved cost", "Estimated saved duration"],
        [
            [
                row["policy"].replace("_", " "),
                f"{num(row['conservative_triggers'])} ({pct_raw(row['conservative_trigger_rate_pct'])})",
                num(row["events_saved"]),
                money(row["estimated_cost_saved_usd"], 2),
                hours(row["estimated_duration_saved_seconds"]),
            ]
            for row in policies
        ],
        [2600, 1900, 1500, 1700, 1800],
    )
    add_para(doc, "Interpretation: stopping policy is a concrete efficiency lever. The best single policy in this run was same_result_no_progress_burst, saving about $1,001 conservatively.")


def q6_same_model(doc: Document, data: dict) -> None:
    rows = data["terminal_gpt5_scaffolds"]
    add_heading(doc, "6. Does The Same Model Behave Differently Under Different Scaffolds?", 1)
    add_callout(
        doc,
        "Direct answer",
        "Yes. In Terminal-Bench trajectory runs for gpt-5@openai, pass rate ranged from 53.1% under codex to about 35.0% under mini-swe-agent, with large duration differences across scaffolds.",
    )
    add_figure(doc, data["figures"]["scaffold"], "Figure 7. Same model, different scaffold outcomes for gpt-5@openai.")
    add_table(
        doc,
        ["Agent scaffold", "Runs", "Pass rate", "Median duration", "Median steps", "Median tool calls"],
        [
            [
                row["agent"],
                num(row["runs"]),
                pct(row["pass_rate"]),
                f"{num(row['median_duration_seconds'])} sec",
                num(row["median_steps"]),
                num(row["median_tool_calls"]),
            ]
            for row in rows
        ],
        [1800, 1300, 1400, 1800, 1500, 1500],
    )
    add_para(
        doc,
        "Interpretation: model-only benchmarking is incomplete. The harness, tools, context strategy, retry behavior, and verification loop can materially change the outcome for the same model.",
    )


def q7_failure_patterns(doc: Document, data: dict) -> None:
    rows = data["terminal_failure_patterns"]
    tool = next(row for row in rows if row["rule_based_rca"] == "TOOL_EXECUTION_FAILURE")
    add_heading(doc, "7. What Are The Most Expensive Failure Patterns?", 1)
    add_callout(
        doc,
        "Direct answer",
        "By estimated avoidable waste, possible missing verification and tool execution failure are the top Terminal-Bench patterns. By total observed spend trapped in failed runs, tool execution failure is much larger, with about $9,522 in observed cost.",
    )
    add_figure(doc, data["figures"]["failure_patterns"], "Figure 8. Estimated avoidable Terminal-Bench cost by failure pattern.")
    add_table(
        doc,
        ["Failure pattern", "Runs", "Failed runs", "Estimated avoidable cost", "Observed cost in pattern"],
        [
            [
                row["rule_based_rca"].replace("_", " ").title(),
                num(row["runs"]),
                num(row["failed_runs"]),
                money(row["estimated_wasted_cost_usd"], 2),
                money(row["total_observed_cost_usd"], 0),
            ]
            for row in rows
            if row["rule_based_rca"] != "MISSING_TRAJECTORY"
            and row["estimated_wasted_cost_usd"] is not None
        ],
        [2700, 1200, 1300, 2100, 2000],
    )
    add_para(
        doc,
        f"Interpretation: teams should prioritize failure patterns by economic impact, not only by count. Tool execution failure has {num(tool['runs'])} runs and {money(tool['total_observed_cost_usd'])} observed cost in this analysis.",
    )


def q8_pass_rate(doc: Document, data: dict) -> None:
    t = data["terminal_overview"]
    s = data["swe_overview_direct"]
    sq = data["swe_quality"]
    add_heading(doc, "8. Why Is Pass Rate Alone A Weak Agent Benchmark?", 1)
    add_callout(
        doc,
        "Direct answer",
        f"Terminal-Bench pass rate is {pct(t['pass_rate'])}, but only {pct(t['trajectory_coverage'])} of runs have usable trajectories. SWE-agent pass rate is {pct(s['pass_rate'])} on {num(s['runs'])} full processed runs, but cost and duration are unavailable. Pass rate answers whether a run worked; it does not answer why, at what cost, or what system behavior caused the result.",
    )
    add_metric_table(
        doc,
        [
            ["Outcome", "Pass/fail", "Useful but too coarse for system improvement."],
            ["Observability", f"Terminal-Bench trajectory coverage {pct(t['trajectory_coverage'])}; SWE-agent {pct(s['trajectory_coverage'])}", "A run without a trajectory cannot support detailed RCA."],
            ["FinOps", "Terminal-Bench cost available; SWE-agent cost unavailable", "Cost-per-success claims need measured cost fields."],
            ["Root cause", "Tool errors, missing verification, repeats, late failures", "These require step-level signals, not only final reward."],
            ["Cutoff and metadata", f"SWE-agent source last modified {sq['source_last_modified']}; snapshot cutoff unavailable", "Benchmark claims should disclose coverage and missing metadata."],
        ],
    )
    add_para(
        doc,
        "Interpretation: stronger agent evaluation combines pass rate with trajectory coverage, cost per success, wasted actions, error events, repeat behavior, verification behavior, and root-cause classification.",
    )


def afwb_section(doc: Document, data: dict) -> None:
    totals = data["afwb_totals"]
    rows = data["afwb_scenarios"]
    add_heading(doc, "Controlled RCA Proof: AFWB Foundation", 1)
    add_callout(
        doc,
        "Why this matters",
        f"AFWB Foundation is not a public-scale benchmark. It is a deterministic proof layer with {totals['scenarios']} scenarios that name the root cause, first observable failure, remediation, and avoidable usage.",
        fill=GOLD_FILL,
    )
    add_table(
        doc,
        ["Scenario", "Root cause", "Remediation", "Avoidable usage"],
        [
            [
                row["scenario_id"],
                row["root_cause_category"],
                row["remediation"]["remediation_type"],
                f"{row['avoidable_usage']['tool_calls']} tool calls; {row['avoidable_usage']['latency_ms']} ms; {row['avoidable_usage']['input_tokens'] + row['avoidable_usage']['output_tokens']} tokens",
            ]
            for row in rows
        ],
        [2300, 2500, 2500, 2100],
    )
    add_metric_table(
        doc,
        [
            ["Avoidable tool calls", num(totals["avoidable_tool_calls"]), "Shows excess tool usage in controlled failures."],
            ["Avoidable latency", f"{num(totals['avoidable_latency_ms'])} ms", "Measures delay from preventable behavior."],
            ["Avoidable tokens", num(totals["avoidable_input_tokens"] + totals["avoidable_output_tokens"]), "Captures prompt/output overhead in controlled failures."],
        ],
        fill=GOLD_FILL,
    )


def interpretation_section(doc: Document, data: dict) -> None:
    add_heading(doc, "What These Benchmarks Let Teams Ask", 1)
    add_para(
        doc,
        "The common thread across the datasets is not a single universal ranking. It is a set of operational questions that pass rate alone cannot answer.",
    )
    add_table(
        doc,
        ["Operational question", "Benchmark evidence", "Why the answer matters"],
        [
            ["Where is budget being wasted?", "Terminal-Bench cost, duration, and waste estimates", "Prioritize failures by dollars and time."],
            ["Which failures are preventable?", "AFWB root-cause and remediation labels", "Separate model limits from system design issues."],
            ["Which runs should stop earlier?", "Terminal-Bench early-stop simulations", "Reduce cost and delay from hopeless loops."],
            ["Which successful runs are fragile?", "Successful-run errors, repeats, and wasted actions", "Improve reliability even when pass rate looks acceptable."],
            ["Is the model or scaffold responsible?", "Same-model, different-scaffold comparison", "Tune harness and tool interface before defaulting to model replacement."],
            ["Can we trust the benchmark row?", "Trajectory, cost, duration, and cutoff coverage", "Avoid making claims from missing or incompatible fields."],
        ],
        [2700, 3000, 3300],
    )
    add_callout(
        doc,
        "Do not combine these into one universal leaderboard",
        "Terminal-Bench, SWE-agent, and AFWB answer different questions. A credible evaluation keeps benchmark families separate and compares systems only within compatible measurement contexts.",
        fill=GOLD_FILL,
    )


def lumniverse_section(doc: Document, data: dict) -> None:
    add_heading(doc, "How Lumniverse Helps", 1)
    add_callout(
        doc,
        "Solution thesis",
        "Lumniverse turns agent trajectories into operational intelligence: why agents fail, where they waste money and time, which scaffold behaviors drive outcomes, and what teams should change next.",
        fill=TEAL_FILL,
    )
    add_table(
        doc,
        ["Problem exposed by the benchmarks", "Lumniverse capability", "Practical outcome"],
        [
            ["Failed runs burn budget before failing", "Cost and waste attribution by run, task, model, and scaffold", "Teams can quantify the economic impact of failure."],
            ["Successful runs hide waste", "Trajectory-level efficiency scoring", "Teams can optimize passing runs, not only failing ones."],
            ["Agents repeat failed actions", "Loop, retry, and same-result detection", "Teams can set better stop, retry, and escalation policies."],
            ["Model-only benchmarks mislead", "Model + scaffold + tool comparison", "Teams can improve harness design instead of only swapping models."],
            ["Root cause is unclear", "RCA for tool, context, verification, retry, and submit failures", "Teams get actionable explanations, not just aggregate scores."],
            ["Benchmark findings do not automatically transfer to production", "Benchmark-to-production observability layer", "Teams can monitor the same failure modes in live agent systems."],
        ],
        [3100, 3000, 3100],
        header_fill=TEAL_FILL,
    )
    add_para(
        doc,
        "The product value is deliberately downstream of the evidence: benchmarks reveal that pass rate is insufficient; Lumniverse gives teams the trajectory intelligence needed to act on that reality.",
        bold=True,
        color=NAVY,
    )


def build_doc() -> None:
    if not DATA_PATH.exists():
        raise SystemExit(f"Missing {DATA_PATH}. Run scripts/prepare_lumniverse_report_assets.py first.")
    data = json.loads(DATA_PATH.read_text())
    doc = Document()
    configure_document(doc)

    title_page(doc, data)
    executive_summary(doc, data)
    methodology(doc, data)
    doc.add_page_break()
    q1_cost_failed_runs(doc, data)
    doc.add_page_break()
    q2_successful_waste(doc, data)
    doc.add_page_break()
    q3_repeated_actions(doc, data)
    doc.add_page_break()
    q4_divergence(doc, data)
    doc.add_page_break()
    q5_early_stop(doc, data)
    doc.add_page_break()
    q6_same_model(doc, data)
    doc.add_page_break()
    q7_failure_patterns(doc, data)
    doc.add_page_break()
    q8_pass_rate(doc, data)
    doc.add_page_break()
    afwb_section(doc, data)
    doc.add_page_break()
    interpretation_section(doc, data)
    doc.add_section(WD_SECTION.NEW_PAGE)
    lumniverse_section(doc, data)

    OUT_DOCX.parent.mkdir(exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"Wrote {OUT_DOCX.relative_to(ROOT)}")


if __name__ == "__main__":
    build_doc()
